#!/usr/bin/env python3
"""Generate tests/selfhost/fixtures/link/fmt_docx/src/main.e (L146): `e.fmt.docx` document generation.

Oracles, independent of the Neper code:
  * the expected package bytes come from a second writer here (hand-built WordprocessingML and zipfile in
    ZIP_STORED with a fixed date); the static parts (styles, relationships boilerplate) are read out of the
    module's source, the dynamic ones (document, numbering, relationships, content types, core) are rebuilt;
  * python-docx must open the bytes and report the same paragraphs, styles, table cells, inline images,
    core properties and page size, and lxml reads the paragraphs for the `e.fmt.ooxml.paragraphs` check;
  * the fixture reads the file back through `e.fmt.ooxml.paragraphs`.
"""
import io
import re
import struct
import sys
import zipfile
import zlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lane_e_common import ROOT, Checks, lit, write_fixture  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
SOURCE = (ROOT / "lib/e/fmt/docx.e").read_text(encoding="utf-8")


def raw_of(fn):
    m = re.search(r"fn %s\(\) -> str \{\n    ret r#\"(.*?)\"#\n\}" % fn, SOURCE, re.S)
    return m.group(1)


def esc(s, attr=False):
    s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return s.replace('"', "&quot;") if attr else s


STYLE_IDS = {"Normal": "", "Title": "Title", "Heading1": "Heading1", "Heading2": "Heading2", "Heading3": "Heading3"}


class R:
    def __init__(self, text, bold=False, italic=False, underline=False, half=0, color=""):
        self.text, self.bold, self.italic, self.underline, self.half, self.color = text, bold, italic, underline, half, color

    def xml(self):
        s = "<w:r>"
        if self.bold or self.italic or self.underline or self.half or self.color:
            s += "<w:rPr>"
            if self.bold:
                s += "<w:b/>"
            if self.italic:
                s += "<w:i/>"
            if self.color:
                s += '<w:color w:val="%s"/>' % self.color
            if self.half:
                s += '<w:sz w:val="%d"/>' % self.half
            if self.underline:
                s += '<w:u w:val="single"/>'
            s += "</w:rPr>"
        buf = ""
        for ch in self.text:
            if ch in "\t\n":
                if buf:
                    s += '<w:t xml:space="preserve">%s</w:t>' % esc(buf)
                    buf = ""
                s += "<w:tab/>" if ch == "\t" else "<w:br/>"
            else:
                buf += ch
        if buf:
            s += '<w:t xml:space="preserve">%s</w:t>' % esc(buf)
        return s + "</w:r>"

    def neper(self):
        return "docx.Run { text: %s, bold: %s, italic: %s, underline: %s, half_points: %du32, color: %s }" % (
            lit(self.text), b(self.bold), b(self.italic), b(self.underline), self.half, lit(self.color))


def b(v):
    return "true" if v else "false"


class P:
    def __init__(self, runs, style="Normal", align="Left", lst="None", list_id=0, level=0):
        self.runs, self.style, self.align, self.lst, self.list_id, self.level = runs, style, align, lst, list_id, level

    def xml(self):
        s = "<w:p>"
        styled = self.style != "Normal"
        listed = self.lst != "None"
        if styled or listed or self.align != "Left":
            s += "<w:pPr>"
            if styled:
                s += '<w:pStyle w:val="%s"/>' % STYLE_IDS[self.style]
            elif listed:
                s += '<w:pStyle w:val="ListParagraph"/>'
            if listed:
                num = 1 if self.lst == "Bullet" else self.list_id + 1
                s += '<w:numPr><w:ilvl w:val="%d"/><w:numId w:val="%d"/></w:numPr>' % (self.level, num)
            if self.align != "Left":
                s += '<w:jc w:val="%s"/>' % {"Center": "center", "Right": "right", "Justify": "both"}[self.align]
            s += "</w:pPr>"
        return s + "".join(r.xml() for r in self.runs) + "</w:p>"

    def neper(self, name):
        runs = "%s_runs" % name
        return runs, "    let %s = [%d]docx.Run{ %s }\n" % (runs, len(self.runs), ", ".join(r.neper() for r in self.runs))

    def expr(self, runs):
        return "docx.Para { style: .%s, align: .%s, list: .%s, list_id: %du32, level: %du32, runs: %s[0..] }" % (
            self.style, self.align, self.lst, self.list_id, self.level, runs)


def png(w, h):
    raw = b"".join(b"\x00" + bytes(sum(([(x * 40 + y * 70) % 256, 90, 200] for x in range(w)), [])) for y in range(h))

    def chunk(kind, data):
        c = struct.pack(">I", len(data)) + kind + data
        return c + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def jpeg(w, h):
    from PIL import Image
    out = io.BytesIO()
    Image.new("RGB", (w, h), (200, 30, 30)).save(out, "JPEG", quality=50)
    return out.getvalue()


PNG = png(5, 3)
JPG = jpeg(8, 6)

TITLE = [R("Quarterly report")]
H1 = [R("Summary & scope")]
MIXED = [R("Plain, "), R("bold", bold=True), R(" italic ", italic=True), R("underlined", underline=True), R(" red 14pt", half=28, color="C00000"),
         R(" a\tb\nc "), R("<tag> & \"quote\"", bold=True, italic=True, underline=True)]
CENTER = [R("Centered line")]
H2 = [R("Details")]
BL0 = [R("first bullet")]
BL1 = [R("nested bullet")]
NL1A = [R("one")]
NL1B = [R("two")]
NL2A = [R("again one")]
NL2B = [R("again two")]
H3 = [R("The end")]


def cell(*paras):
    return list(paras)


TABLE = {
    "widths": [2000, 3000, 4000],
    "header": True,
    "rows": [
        [cell(P([R("Item", bold=True)])), cell(P([R("Qty", bold=True)])), cell(P([R("Notes", bold=True)]))],
        [cell(P([R("Apple")])), cell(P([R("12")], align="Right")), cell(P([R("first note")]), P([R("second note & more")]))],
        [cell(P([R("Pear")])), cell(), cell(P([R("")]))],
        [cell(P([R("Short row")]))],
    ],
}

BLOCKS = [
    ("p", P(TITLE, "Title")),
    ("p", P(H1, "Heading1")),
    ("p", P(MIXED)),
    ("p", P(CENTER, align="Center")),
    ("p", P(H2, "Heading2")),
    ("p", P(BL0, lst="Bullet")),
    ("p", P(BL1, lst="Bullet", level=1)),
    ("p", P(NL1A, lst="Number", list_id=1)),
    ("p", P(NL1B, lst="Number", list_id=1)),
    ("p", P(NL2A, lst="Number", list_id=2)),
    ("p", P(NL2B, lst="Number", list_id=2)),
    ("t", TABLE),
    ("i", {"data": PNG, "cx": 0, "cy": 0, "alt": 'A "pixel" chart', "jpeg": False, "px": (5, 3)}),
    ("i", {"data": JPG, "cx": 1905000, "cy": 1428750, "alt": "", "jpeg": True, "px": (8, 6)}),
    ("b", None),
    ("p", P(H3, "Heading3", align="Justify")),
]
META = {"title": "Report – Q3 & more", "creator": "Neper <docs>", "created": "2024-05-17T10:00:00Z", "page": "A4"}


def para_xml_list(paras):
    return "".join(p.xml() for p in paras)


def build_document(blocks, meta):
    body = ""
    images = []
    numbered = 0
    for kind, v in blocks:
        if kind == "p":
            body += v.xml()
            if v.lst == "Number":
                numbered = max(numbered, v.list_id)
        elif kind == "t":
            t = v
            body += '<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="0" w:type="auto"/></w:tblPr><w:tblGrid>'
            body += "".join('<w:gridCol w:w="%d"/>' % w for w in t["widths"]) + "</w:tblGrid>"
            for ri, row in enumerate(t["rows"]):
                body += "<w:tr>"
                if t["header"] and ri == 0:
                    body += "<w:trPr><w:tblHeader/></w:trPr>"
                for ci, wd in enumerate(t["widths"]):
                    body += '<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/></w:tcPr>' % wd
                    paras = row[ci] if ci < len(row) else []
                    body += "".join(p.xml() for p in paras) if paras else "<w:p/>"
                    for p in paras:
                        if p.lst == "Number":
                            numbered = max(numbered, p.list_id)
                    body += "</w:tc>"
                body += "</w:tr>"
            body += "</w:tbl><w:p/>"
        elif kind == "b":
            body += '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'
        else:
            images.append(v)
            n = len(images)
            ext = "jpeg" if v["jpeg"] else "png"
            cx = v["cx"] or v["px"][0] * 9525
            cy = v["cy"] or v["px"][1] * 9525
            body += ('<w:p><w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="%d" cy="%d"/>'
                     '<wp:docPr id="%d" name="Picture %d" descr="%s"/><wp:cNvGraphicFramePr>'
                     '<a:graphicFrameLocks xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" noChangeAspect="1"/></wp:cNvGraphicFramePr>'
                     '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
                     '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:nvPicPr><pic:cNvPr id="0" name="image%d.%s"/><pic:cNvPicPr/></pic:nvPicPr>'
                     '<pic:blipFill><a:blip r:embed="rId%d"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
                     '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
                     % (cx, cy, n, n, esc(v["alt"], True), n, ext, n + 2, cx, cy))
    size = '<w:pgSz w:w="12240" w:h="15840"/>' if meta["page"] == "Letter" else '<w:pgSz w:w="11906" w:h="16838"/>'
    document = (XML + '<w:document xmlns:w="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
                'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"><w:body>' % W + body +
                '<w:sectPr>' + size + '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="708" w:footer="708" w:gutter="0"/></w:sectPr></w:body></w:document>')
    types = XML + ('<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>')
    if any(not i["jpeg"] for i in images):
        types += '<Default Extension="png" ContentType="image/png"/>'
    if any(i["jpeg"] for i in images):
        types += '<Default Extension="jpeg" ContentType="image/jpeg"/>'
    types += ('<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
              '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
              '<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>'
              '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/></Types>')
    rels = XML + ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                  '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                  '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/>')
    for n, im in enumerate(images, 1):
        rels += ('<Relationship Id="rId%d" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image%d.%s"/>'
                 % (n + 2, n, "jpeg" if im["jpeg"] else "png"))
    rels += "</Relationships>"
    numbering = XML + '<w:numbering xmlns:w="%s">' % W
    for kind in range(2):
        numbering += '<w:abstractNum w:abstractNumId="%d"><w:multiLevelType w:val="hybridMultilevel"/>' % kind
        for level in range(9):
            numbering += '<w:lvl w:ilvl="%d"><w:start w:val="1"/>' % level
            if kind == 0:
                numbering += '<w:numFmt w:val="bullet"/><w:lvlText w:val="•"/>'
            else:
                numbering += '<w:numFmt w:val="decimal"/><w:lvlText w:val="%%%d."/>' % (level + 1)
            numbering += '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="%d" w:hanging="360"/></w:pPr></w:lvl>' % (720 * (level + 1))
        numbering += "</w:abstractNum>"
    numbering += '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
    for k in range(1, numbered + 1):
        numbering += '<w:num w:numId="%d"><w:abstractNumId w:val="1"/><w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>' % (k + 1)
    numbering += "</w:numbering>"
    core = XML + ('<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" '
                  'xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">')
    if meta["title"]:
        core += "<dc:title>%s</dc:title>" % esc(meta["title"])
    if meta["creator"]:
        core += "<dc:creator>%s</dc:creator>" % esc(meta["creator"])
    if meta["created"]:
        core += '<dcterms:created xsi:type="dcterms:W3CDTF">%s</dcterms:created>' % esc(meta["created"])
    core += "</cp:coreProperties>"
    parts = [("[Content_Types].xml", types.encode()), ("_rels/.rels", raw_of("root_rels").encode()), ("docProps/core.xml", core.encode()),
             ("word/document.xml", document.encode()), ("word/_rels/document.xml.rels", rels.encode()),
             ("word/styles.xml", raw_of("styles_xml").encode()), ("word/numbering.xml", numbering.encode())]
    for n, im in enumerate(images, 1):
        parts.append(("word/media/image%d.%s" % (n, "jpeg" if im["jpeg"] else "png"), im["data"]))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name, data in parts:
            zi = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_STORED
            zi.create_system = 0
            z.writestr(zi, data)
    return out.getvalue(), document


def expected_paragraphs(document_xml):
    from lxml import etree
    ns = {"w": W}
    root = etree.fromstring(document_xml.encode("utf-8"))
    out = []
    for p in root.iter("{%s}p" % W):
        text = ""
        for el in p.iter():
            if el.tag == "{%s}t" % W:
                text += el.text or ""
            elif el.tag == "{%s}tab" % W:
                text += "\t"
            elif el.tag == "{%s}br" % W:
                # e.fmt.ooxml reads every break, a page break included, as a newline.
                text += "\n"
        style = ""
        ps = p.find("w:pPr/w:pStyle", ns)
        if ps is not None:
            style = ps.get("{%s}val" % W)
        out.append((text, style, p.find("w:pPr/w:numPr", ns) is not None))
    return out


def verify(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    assert d.core_properties.title == META["title"], d.core_properties.title
    assert d.core_properties.author == META["creator"]
    assert d.sections[0].page_width == 11906 * 635 and d.sections[0].page_height == 16838 * 635
    styles = [p.style.name for p in d.paragraphs]
    texts = [p.text for p in d.paragraphs]
    assert "Title" in styles and "Heading 1" in styles and "Heading 2" in styles and "Heading 3" in styles and "List Paragraph" in styles
    assert "Quarterly report" in texts and "Summary & scope" in texts
    mixed = [p for p in d.paragraphs if p.text.startswith("Plain, ")][0]
    assert mixed.text == "Plain, bold italic underlined red 14pt a\tb\nc <tag> & \"quote\"", repr(mixed.text)
    bold = [r.text for r in mixed.runs if r.bold]
    assert bold[0] == "bold", bold
    assert [r for r in mixed.runs if r.text == " red 14pt"][0].font.size.pt == 14
    assert str([r for r in mixed.runs if r.text == " red 14pt"][0].font.color.rgb) == "C00000"
    t = d.tables[0]
    assert [[c.text for c in row.cells] for row in t.rows][0] == ["Item", "Qty", "Notes"]
    assert t.rows[1].cells[2].text == "first note\nsecond note & more"
    assert t.style.name == "Table Grid"
    assert len(d.inline_shapes) == 2
    assert d.inline_shapes[0].width == 5 * 9525 and d.inline_shapes[1].width == 1905000
    assert [p.alignment for p in d.paragraphs if p.text == "Centered line"][0] == 1


def main():
    data, document = build_document(BLOCKS, META)
    verify(data)
    empty, empty_document = build_document([], {"title": "", "creator": "", "created": "", "page": "Letter"})
    paras = expected_paragraphs(document)

    c = Checks()
    body = ""
    names = []
    # Model arrays for the full document.
    block_exprs = []
    pn = 0

    def para_block(p):
        nonlocal pn, body
        pn += 1
        name = "p%d" % pn
        runs, text = p.neper(name)
        body += text
        return "docx.paragraph(%s)" % p.expr(runs)

    for kind, v in BLOCKS:
        if kind == "p":
            block_exprs.append(para_block(v))
        elif kind == "t":
            pn += 1
            tname = "t%d" % pn
            cells_by_row = []
            for ri, row in enumerate(v["rows"]):
                cell_names = []
                for ci, paras_ in enumerate(row):
                    paras_exprs = []
                    for pi, p in enumerate(paras_):
                        pname = "%s_r%dc%dp%d" % (tname, ri, ci, pi)
                        runs, text = p.neper(pname)
                        body += text
                        paras_exprs.append(p.expr(runs))
                    cname = "%s_r%dc%d" % (tname, ri, ci)
                    if paras_exprs:
                        body += "    let %s_paras = [%d]docx.Para{ %s }\n" % (cname, len(paras_exprs), ", ".join(paras_exprs))
                        cell_names.append("docx.Cell { paras: %s_paras[0..] }" % cname)
                    else:
                        cell_names.append("docx.Cell { paras: zero }")
                rname = "%s_r%d" % (tname, ri)
                body += "    let %s_cells = [%d]docx.Cell{ %s }\n" % (rname, len(cell_names), ", ".join(cell_names))
                cells_by_row.append("docx.Row { cells: %s_cells[0..] }" % rname)
            body += "    let %s_rows = [%d]docx.Row{ %s }\n" % (tname, len(cells_by_row), ", ".join(cells_by_row))
            body += "    let %s_widths = [%d]u32{ %s }\n" % (tname, len(v["widths"]), ", ".join("%du32" % w for w in v["widths"]))
            block_exprs.append("docx.table(docx.Table { widths: %s_widths[0..], rows: %s_rows[0..], header: %s })" % (tname, tname, b(v["header"])))
        elif kind == "i":
            pn += 1
            body += "    let img%d = %s\n" % (pn, lit(v["data"]))
            block_exprs.append("docx.image(docx.Image { data: img%d, width_emu: %du32, height_emu: %du32, alt: %s })" % (pn, v["cx"], v["cy"], lit(v["alt"])))
        else:
            block_exprs.append("docx.page_break()")
    body += "    let blocks = [%d]docx.Block{ %s }\n" % (len(block_exprs), ", ".join(block_exprs))
    body += "    let meta = docx.Meta { title: %s, creator: %s, created: %s, page: .A4 }\n" % (lit(META["title"]), lit(META["creator"]), lit(META["created"]))
    body += "    let (bytes, bytes_error) = docx.encode(a, blocks[0..], meta)\n"
    c.raw(body)
    c.add("bytes_error != ok")
    c.add("bytes.len != %dusize || !same(bytes, %s)" % (len(data), lit(data)))
    c.raw("    let (read, read_error) = ooxml.paragraphs(a, bytes)\n")
    c.add("read_error != ok")
    c.add("read.len != %dusize" % len(paras))
    for i, (text, style, listed) in enumerate(paras):
        c.add("!str.eq(read[%d].text, %s) || !str.eq(read[%d].style, %s) || %sread[%d].list" % (i, lit(text), i, lit(style), "!" if listed else "", i))
    # An empty document, Letter paper.
    c.raw("    let no_blocks = [1]docx.Block{ docx.page_break() }\n")
    c.raw("    let (blank, blank_error) = docx.encode(a, no_blocks[0..0], docx.Meta { title: \"\", creator: \"\", created: \"\", page: .Letter })\n")
    c.add("blank_error != ok || blank.len != %dusize || !same(blank, %s)" % (len(empty), lit(empty)))
    # Refusals.
    def refuse(n, runs_expr, error, extra=""):
        c.raw("    let bad%d = [1]docx.Run{ %s }\n" % (n, runs_expr))
        c.raw("    let bad%d_block = [1]docx.Block{ docx.paragraph(docx.para(.Normal, bad%d[0..])) }\n" % (n, n))
        c.raw("    let (junk%d, junk%d_error) = docx.encode(a, bad%d_block[0..], docx.default_meta())\n" % (n, n, n))
        c.add("junk%d_error != %s" % (n, error))

    refuse(1, 'docx.run("a\\x01b")', "docx.BadText")
    refuse(2, 'docx.Run { text: "x", bold: false, italic: false, underline: false, half_points: 0u32, color: "red" }', "docx.BadColor")
    refuse(3, 'docx.Run { text: "x", bold: false, italic: false, underline: false, half_points: 2000u32, color: "" }', "docx.TooLarge")
    c.raw("    let tw = [1]u32{ 100u32 }\n")
    c.raw("    let tcell = [1]docx.Cell{ docx.Cell { paras: zero } }\n")
    c.raw("    let trow = [1]docx.Row{ docx.Row { cells: tcell[0..] } }\n")
    c.raw("    let no_widths = [1]docx.Block{ docx.table(docx.Table { widths: zero, rows: trow[0..], header: false }) }\n")
    c.raw("    let (nw, nw_error) = docx.encode(a, no_widths[0..], docx.default_meta())\n")
    c.add("nw_error != docx.BadTable")
    c.raw("    let wide_cells = [2]docx.Cell{ docx.Cell { paras: zero }, docx.Cell { paras: zero } }\n")
    c.raw("    let wide_row = [1]docx.Row{ docx.Row { cells: wide_cells[0..] } }\n")
    c.raw("    let too_wide = [1]docx.Block{ docx.table(docx.Table { widths: tw[0..], rows: wide_row[0..], header: false }) }\n")
    c.raw("    let (tl, tl_error) = docx.encode(a, too_wide[0..], docx.default_meta())\n")
    c.add("tl_error != docx.BadTable")
    c.raw("    let not_image = [1]docx.Block{ docx.image(docx.Image { data: \"GIF89a not a png\", width_emu: 0u32, height_emu: 0u32, alt: \"\" }) }\n")
    c.raw("    let (ni, ni_error) = docx.encode(a, not_image[0..], docx.default_meta())\n")
    c.add("ni_error != docx.BadImage")
    c.raw("    let (bm, bm_error) = docx.encode(a, no_blocks[0..0], docx.Meta { title: \"a\\x01\", creator: \"\", created: \"\", page: .A4 })\n")
    c.add("bm_error != docx.BadMeta")
    header = ("// `e.fmt.docx`: a document with a title, headings, mixed runs (bold, italic, underline, size, colour, a tab and a line\n"
              "// break, markup characters), a centred paragraph, bullet and two restarting numbered lists, a table with a header\n"
              "// row, wrapped cells and a short row, a PNG and a JPEG, a page break and Meta compared byte for byte with an\n"
              "// independent writer; python-docx opened the file when the generator ran; the paragraphs read back through\n"
              "// `e.fmt.ooxml.paragraphs`; and the refusals. Written by scripts/gen_docx_test.py. Every check has its own exit code.\n")
    uses = ["e.os", "e.mem", "e.str", "e.fmt.ooxml as ooxml", "e.fmt.docx as docx"]
    fn_head = ("fn same(x: []const u8, y: []const u8) -> bool {\n    if x.len != y.len { ret false }\n    var i = 0usize\n"
               "    while i < x.len {\n        if x[i] != y[i] { ret false }\n        i += 1usize\n    }\n    ret true\n}\n\n"
               "fn main(a: *mem.Arena, args: []str) -> err {\n")
    target = write_fixture("fmt_docx", header, fn_head + c.text() + "    ret ok\n}\n", uses)
    print("wrote", target, "checks", c.code, "bytes", len(data), "paragraphs", len(paras))


if __name__ == "__main__":
    main()

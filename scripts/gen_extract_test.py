#!/usr/bin/env python3
"""Generate tests/selfhost/fixtures/link/fmt_extract/src/main.e (L153): `e.fmt.extract` detect and extract.

The corpus is written by independent producers -- reportlab (a compressed two-page PDF), python-docx,
python-pptx, openpyxl, hand-built OpenDocument packages, Python's email package (multipart/alternative with
quoted-printable and an RFC 2047 subject, base64 Latin-1, HTML-only), plus text, HTML, XML, JSON, Markdown,
CSV, TSV and vCard files -- and one PDF written by the Neper writer's own independent generator. The
expected text of each is derived without the Neper code: lxml walks for the Office and OpenDocument parts,
openpyxl read-back for the workbook, bs4 and markdown-it for the HTML and Markdown checks, and the
strings that went into the PDF, mail and vCard files.
"""
import base64
import email.message
import email.policy
import io
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lane_e_common import Checks, lit, write_fixture  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
SAMPLES = []  # (var, file name, bytes, format, text, title, author, parts or None)


def add(var, name, data, fmt, text, title="", author="", parts=None, exact=True, contains=None):
    if isinstance(data, str):
        data = data.encode("utf-8")
    SAMPLES.append(dict(var=var, name=name, data=data, fmt=fmt, text=text, title=title, author=author, parts=parts, exact=exact,
                        contains=contains or []))


# ---- plain text -----------------------------------------------------------------------------
add("txt", "notes.txt", b"\xef\xbb\xbfFirst line caf\xc3\xa9\nSecond line \xe2\x82\xac5\n", "Text", "First line café\nSecond line €5\n")

# ---- html -----------------------------------------------------------------------------------
HTML = ("<!DOCTYPE html><html><head><title> Neper &amp; friends </title><style>p{color:red}</style></head>"
        "<body><h1>Heading One</h1><p>First   paragraph with <b>bold</b> and <a href=\"x\">a link</a>.</p>"
        "<script>var hidden = 1;</script><ul><li>alpha</li><li>beta &amp; gamma</li></ul>"
        "<table><tr><th>Name</th><th>Qty</th></tr><tr><td>Apple</td><td>3</td></tr></table>"
        "<!-- a comment --><p>caf&eacute; &euro;5</p></body></html>")
HTML_TEXT = "Heading One\nFirst paragraph with bold and a link.\nalpha\nbeta & gamma\nName\tQty\nApple\t3\ncafé €5"
add("html", "page.html", HTML, "Html", HTML_TEXT, title="Neper & friends")
add("html_in_txt", "page.txt", HTML, "Html", HTML_TEXT, title="Neper & friends")

# ---- xml, json ------------------------------------------------------------------------------
add("xml", "data.xml", '<?xml version="1.0"?><root><item id="1">One</item><item>Two &amp; three</item><!-- c --><note>  spaced  </note></root>',
    "Xml", "One\nTwo & three\nspaced", parts=3)
JSON = '{"title":"Neper","tags":["a","b"],"n":3,"nested":{"k":"v"},"flag":true,"nothing":null}'
add("json", "data.json", JSON, "Json", "title\nNeper\ntags\na\nb\nn\nnested\nk\nv\nflag\nnothing", parts=11)
add("json_bare", "download", JSON, "Json", "title\nNeper\ntags\na\nb\nn\nnested\nk\nv\nflag\nnothing", parts=11)

# ---- markdown (checked against markdown-it) ---------------------------------------------------
MD = ("# Neper Guide\n\nSome *emphasis* and `code` here, a [link](http://x.y/z) and ![logo](a.png).\n\n"
      "- item one\n- item **two**\n\n> quoted text\n\n1. first\n2. second\n\n```\ncode line\n```\n\n---\n\nTail paragraph.\n")
MD_TEXT = ("Neper Guide\n\nSome emphasis and code here, a link and logo.\n\nitem one\nitem two\n\nquoted text\n\nfirst\nsecond\n\ncode line\n\n\nTail paragraph.\n")
add("md", "guide.md", MD, "Markdown", MD_TEXT, title="Neper Guide")

# ---- csv, tsv -------------------------------------------------------------------------------
add("csv", "table.csv", 'name,qty,note\n"Smith, J",2,"line1\nline2"\n"quote ""q""",3,\n', "Csv",
    'name\tqty\tnote\nSmith, J\t2\tline1\nline2\nquote "q"\t3\t', parts=3)
add("tsv", "table.tsv", "a\tb\n1\t2\n", "Csv", "a\tb\n1\t2", parts=2)

# ---- vcard ----------------------------------------------------------------------------------
VCF = ("BEGIN:VCARD\r\nVERSION:3.0\r\nFN:Ada Lovelace\r\nN:Lovelace;Ada;;;\r\nORG:Analytical Engines\r\nTEL;TYPE=WORK:+44 1 2345\r\nEMAIL:ada@example.org\r\nEND:VCARD\r\n"
       "BEGIN:VCARD\r\nVERSION:3.0\r\nFN:Grace Hopper\r\nN:Hopper;Grace;;;\r\nEMAIL:grace@example.org\r\nEND:VCARD\r\n")
add("vcf", "people.vcf", VCF, "VCard", None, parts=2, exact=False,
    contains=["Ada Lovelace", "Analytical Engines", "+44 1 2345", "ada@example.org", "Grace Hopper", "grace@example.org"])


# ---- pdf: reportlab, compressed -------------------------------------------------------------
def make_pdf():
    """A reportlab PDF whose content streams pypdf then compresses with FlateDecode alone (reportlab's own
    page compression is ASCII85 plus Flate, which e.fmt.pdf does not read)."""
    from reportlab.pdfgen import canvas
    import pypdf
    out = io.BytesIO()
    c = canvas.Canvas(out, pageCompression=0, invariant=1)
    c.setTitle("Quarterly summary report")
    c.setAuthor("Ada Lovelace")
    c.setFont("Helvetica", 12)
    c.drawString(72, 750, "Quarterly summary")
    c.drawString(72, 730, "Revenue grew 12% (net) in Q3")
    c.showPage()
    c.setFont("Helvetica-Bold", 12)
    c.drawString(72, 750, "Appendix A")
    c.showPage()
    c.save()
    reader = pypdf.PdfReader(io.BytesIO(out.getvalue()))
    writer = pypdf.PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    for page in writer.pages:
        page.compress_content_streams()
    writer.add_metadata({"/Title": "Quarterly summary report", "/Author": "Ada Lovelace"})
    final = io.BytesIO()
    writer.write(final)
    data = final.getvalue()
    assert b"/FlateDecode" in data
    return data


add("pdf", "report.pdf", make_pdf(), "Pdf", None, title="Quarterly summary report", author="Ada Lovelace", parts=2, exact=False,
    contains=["Quarterly summary", "Revenue grew 12% (net) in Q3", "Appendix A"])


def make_neper_pdf():
    import gen_pdf_write_test as g
    mini_bytes, mini_neper, full, full_neper, doc = g.build_main()
    return full


add("pdf_neper", "written.pdf", make_neper_pdf(), "Pdf", None, title="Quarterly report – Q3 café", parts=2, exact=False,
    contains=["Apple pie", "A new paragraph after a forced break", "Page two"])


# ---- docx: python-docx; expected by an lxml walk over the main part ------------------------------
def walk_paragraphs(xml_bytes, ns_uri):
    from lxml import etree
    root = etree.fromstring(xml_bytes)
    out = []
    for p in root.iter("{%s}p" % ns_uri):
        text = ""
        for el in p.iter():
            local = etree.QName(el).localname
            if local == "t" and etree.QName(el).namespace == ns_uri:
                text += el.text or ""
            elif local == "tab" and etree.QName(el).namespace == ns_uri:
                text += "\t"
            elif local == "br" and etree.QName(el).namespace == ns_uri:
                text += "\n"
        out.append(text)
    return out


def make_docx():
    import docx
    d = docx.Document()
    d.core_properties.title = "Memo – Q3"
    d.core_properties.author = "Ada Lovelace"
    d.add_heading("Findings", level=1)
    d.add_paragraph("Revenue grew twelve percent.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Region"
    t.cell(0, 1).text = "Growth"
    t.cell(1, 0).text = "North"
    t.cell(1, 1).text = "12%"
    d.add_paragraph("Closing remark with a tab\tand more.")
    out = io.BytesIO()
    d.save(out)
    data = out.getvalue()
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")
    return data, "\n".join(walk_paragraphs(xml, W)), len(walk_paragraphs(xml, W))


data, text, count = make_docx()
add("docx", "memo.docx", data, "Docx", text, title="Memo – Q3", author="Ada Lovelace", parts=count)


# ---- pptx: python-pptx --------------------------------------------------------------------------
def make_pptx():
    from pptx import Presentation
    p = Presentation()
    s1 = p.slides.add_slide(p.slide_layouts[0])
    s1.shapes.title.text = "Deck title"
    s1.placeholders[1].text = "A subtitle"
    s2 = p.slides.add_slide(p.slide_layouts[1])
    s2.shapes.title.text = "Second slide"
    s2.placeholders[1].text = "bullet one\nbullet two"
    p.core_properties.title = "Quarterly deck"
    p.core_properties.author = "Grace Hopper"
    out = io.BytesIO()
    p.save(out)
    data = out.getvalue()
    z = zipfile.ZipFile(io.BytesIO(data))
    texts = []
    for n in range(1, 3):
        paras = walk_paragraphs(z.read("ppt/slides/slide%d.xml" % n), A)
        texts.append("\n".join(paras))
    return data, "\n\n".join(texts)


data, text = make_pptx()
add("pptx", "deck.pptx", data, "Pptx", text, title="Quarterly deck", author="Grace Hopper", parts=2)


# ---- xlsx: openpyxl ---------------------------------------------------------------------------
def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alpha"
    rows = [["x", "y"], [1, 2.5], ["t", True]]
    for r in rows:
        ws.append(r)
    wb.create_sheet("Beta").append(["only"])
    out = io.BytesIO()
    wb.save(out)
    data = out.getvalue()
    back = openpyxl.load_workbook(io.BytesIO(data))
    lines = []
    for sheet in back.worksheets:
        lines.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            cells = []
            for v in row:
                if isinstance(v, bool):
                    cells.append("true" if v else "false")
                elif isinstance(v, float) and v == int(v):
                    cells.append(str(int(v)))
                elif v is None:
                    cells.append("")
                else:
                    cells.append(str(v))
            lines.append("\t".join(cells))
    return data, "\n".join(lines)


data, text = make_xlsx()
add("xlsx", "sheet.xlsx", data, "Xlsx", text, parts=2)


# ---- odf: hand-built packages ---------------------------------------------------------------------
OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
DRAW = "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
META = ('<office:document-meta xmlns:office="%s" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0">'
        '<office:meta><dc:title>%%s</dc:title><dc:creator>%%s</dc:creator></office:meta></office:document-meta>' % OFFICE)


def odf(mimetype, body, title, creator):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), mimetype, compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", ('<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="%s" xmlns:text="%s" xmlns:draw="%s" xmlns:table="%s">'
                                   % (OFFICE, TEXT, DRAW, TABLE)) + body + "</office:document-content>", compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("meta.xml", META % (title, creator), compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


def odf_paragraphs(data):
    from lxml import etree
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("content.xml"))
    out = []

    def text_of(e):
        parts = [e.text or ""]
        for child in e:
            q = etree.QName(child)
            if q.localname == "s":
                parts.append(" " * int(child.get("{%s}c" % TEXT, "1")))
            elif q.localname == "tab":
                parts.append("\t")
            elif q.localname == "line-break":
                parts.append("\n")
            else:
                parts.append(text_of(child))
            parts.append(child.tail or "")
        return "".join(parts)

    for el in root.iter("{%s}p" % TEXT, "{%s}h" % TEXT):
        out.append(text_of(el))
    return out


ODT = odf("application/vnd.oasis.opendocument.text",
          '<office:body><office:text><text:h text:outline-level="1">Heading</text:h><text:p>Body with <text:span>span</text:span> and a<text:tab/>tab</text:p>'
          '<text:list><text:list-item><text:p>item</text:p></text:list-item></text:list></office:text></office:body>', "Letter", "Ada Lovelace")
add("odt", "letter.odt", ODT, "Odt", "\n".join(odf_paragraphs(ODT)), title="Letter", author="Ada Lovelace", parts=len(odf_paragraphs(ODT)))
ODS = odf("application/vnd.oasis.opendocument.spreadsheet",
          '<office:body><office:spreadsheet><table:table><table:table-row><table:table-cell><text:p>a1</text:p></table:table-cell>'
          '<table:table-cell><text:p>b1</text:p></table:table-cell></table:table-row></table:table></office:spreadsheet></office:body>', "Calc", "Grace Hopper")
add("ods", "calc.ods", ODS, "Ods", "\n".join(odf_paragraphs(ODS)), title="Calc", author="Grace Hopper", parts=len(odf_paragraphs(ODS)))
ODP = odf("application/vnd.oasis.opendocument.presentation",
          '<office:body><office:presentation><draw:page><draw:frame><draw:text-box><text:p>Slide one</text:p><text:p>sub</text:p></draw:text-box></draw:frame></draw:page>'
          '<draw:page><draw:frame><draw:text-box><text:p>Slide two</text:p></draw:text-box></draw:frame></draw:page></office:presentation></office:body>', "Show", "Ada Lovelace")
add("odp", "show.odp", ODP, "Odp", "Slide one\nsub\n\nSlide two", title="Show", author="Ada Lovelace", parts=2)


# ---- mail: Python's email package ---------------------------------------------------------------
def make_mail_alternative():
    m = email.message.EmailMessage(policy=email.policy.SMTP)
    m["Subject"] = "Réunion café"
    m["From"] = "Ada Löve <ada@example.org>"
    m["To"] = "grace@example.org"
    m.set_content("Bonjour à tous,\nle café est prêt.\n", cte="quoted-printable")
    m.add_alternative("<html><body><p>Bonjour <b>html</b></p></body></html>", subtype="html")
    return m.as_bytes()


mail1 = make_mail_alternative()
add("mail_alt", "message.eml", mail1, "Mail", "Bonjour à tous,\nle café est prêt.", title="Réunion café",
    author="Ada Löve <ada@example.org>", parts=1)

latin = ("From: Grace <grace@example.org>\r\nSubject: Grüße\r\nMIME-Version: 1.0\r\nContent-Type: text/plain; charset=iso-8859-1\r\n"
         "Content-Transfer-Encoding: base64\r\n\r\n" + base64.encodebytes("Grüße aus Köln".encode("latin-1")).decode() + "\r\n")
add("mail_b64", "latin.eml", latin, "Mail", "Grüße aus Köln", title="Grüße", author="Grace <grace@example.org>", parts=1, exact=True)
# The subject here is a bare 8-bit header; the checked title is what the file carries.
add("mail_html", "page.eml", "Subject: Only html\r\nFrom: a@b.c\r\nContent-Type: text/html; charset=utf-8\r\n\r\n<html><body><p>Only <i>html</i> body</p></body></html>\r\n",
    "Mail", "Only html body", title="Only html", author="a@b.c", parts=1)
add("mail_sniffed", "inbox_item", "Received: from x\r\nSubject: Sniffed\r\n\r\nplain body\r\n", "Mail", "plain body", title="Sniffed", parts=1)

# ---- no extension, plain words; unknown binary ----------------------------------------------------
add("words", "README", "just some plain words\nover two lines\n", "Text", "just some plain words\nover two lines\n")
BINARY = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\x00IEND\xaeB`\x82"


def checks():
    from bs4 import BeautifulSoup
    from markdown_it import MarkdownIt
    # Independent sanity checks of the hand-written expectations.
    soup = BeautifulSoup(HTML, "html.parser")
    for s in soup(["script", "style"]):
        s.decompose()
    seen = "".join(soup.get_text("").split())
    for line in HTML_TEXT.replace("\t", " ").split("\n"):
        assert "".join(line.split()) in seen, line
    assert "hidden" not in HTML_TEXT
    md_html = MarkdownIt().render(MD)
    md_text = " ".join(BeautifulSoup(md_html, "html.parser").get_text(" ").split())
    assert " ".join(MD_TEXT.split()) == md_text.replace(" ,", ",").replace(" .", ".") or all(
        w in md_text for w in "Neper Guide emphasis code link item one two quoted text first second code line Tail paragraph".split()), md_text


def main():
    checks()
    c = Checks()
    body = ""
    for s in SAMPLES:
        body += "    let %s = %s\n" % (s["var"], lit(s["data"]))
    c.raw(body)
    for s in SAMPLES:
        v = s["var"]
        c.raw("    let (%s_d, %s_e) = x.extract(a, %s, %s)\n" % (v, v, v, lit(s["name"])))
        c.add("%s_e != ok" % v)
        c.add("%s_d.format != .%s" % (v, s["fmt"]))
        if s["exact"]:
            c.add("!str.eq(%s_d.text, %s)" % (v, lit(s["text"])))
        else:
            for needle in s["contains"]:
                c.add("!str.contains(%s_d.text, %s)" % (v, lit(needle)))
        if s["title"] or s["author"] or s["fmt"] in ("Html", "Markdown", "Mail", "Pdf", "Docx", "Pptx", "Odt", "Odp", "Ods"):
            c.add("!str.eq(%s_d.title, %s) || !str.eq(%s_d.author, %s)" % (v, lit(s["title"]), v, lit(s["author"])))
        if s["parts"] is not None:
            c.add("%s_d.parts != %dusize" % (v, s["parts"]))
    # Detection on its own: format by content, extension as the tie-break, nothing for binary data.
    c.add("x.detect(a, %s, \"\") != .Pdf" % "pdf")
    c.add("x.detect(a, html_in_txt, \"page.txt\") != .Html")
    c.add("x.detect(a, json_bare, \"\") != .Json")
    c.add("x.detect(a, words, \"\") != .Text")
    c.add("x.detect(a, mail_sniffed, \"\") != .Mail")
    c.add("x.detect(a, docx, \"renamed.bin\") != .Docx")
    c.raw("    let binary = %s\n" % lit(BINARY))
    c.add("x.detect(a, binary, \"image.png\") != .Unknown")
    c.raw("    let (bin_d, bin_e) = x.extract(a, binary, \"image.png\")\n")
    c.add("bin_e != x.Unsupported || bin_d.format != .Unknown")
    c.raw("    let (nothing_d, nothing_e) = x.extract(a, \"\", \"\")\n")
    c.add("nothing_e != x.Unsupported")
    c.raw("    let (bad_d, bad_e) = x.extract_as(a, \"not a pdf at all\", .Pdf)\n")
    c.add("bad_e != x.Corrupt")
    c.raw("    let (latin_d, latin_e) = x.extract_as(a, \"caf\\xe9\", .Text)\n")
    c.add("latin_e != x.Encoding")
    c.add("!str.eq(x.media_type(.Docx), \"application/vnd.openxmlformats-officedocument.wordprocessingml.document\")")
    header = ("// `e.fmt.extract`: a corpus written by independent producers (reportlab, python-docx, python-pptx, openpyxl, hand-built\n"
              "// OpenDocument, Python's email package, text, HTML, XML, JSON, Markdown, CSV, TSV, vCard and the Neper PDF writer)\n"
              "// detected and extracted: format, text (exact where the producer fixes it, else the lines it drew), title, author\n"
              "// and part count, then detection by content over extension, binary data refused. Written by\n"
              "// scripts/gen_extract_test.py. Every check has its own exit code.\n")
    uses = ["e.os", "e.mem", "e.str", "e.fmt.extract as x"]
    fn_head = "fn main(a: *mem.Arena, args: []str) -> err {\n"
    target = write_fixture("fmt_extract", header, fn_head + c.text() + "    ret ok\n}\n", uses)
    print("wrote", target, "checks", c.code, "samples", len(SAMPLES))


if __name__ == "__main__":
    main()
